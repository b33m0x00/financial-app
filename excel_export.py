"""Export inputs to an .xlsx where every ratio is a live Excel formula (same logic as the course templates)."""
from __future__ import annotations

import io
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from fin import INPUTS, RATIOS, SCORE_INFO

# Excel formula per ratio. {key} = this year's cell, {key@prev} = previous year's cell.
RATIO_XL = {
    "return_on_sales": "={pretax_profit}/{revenue}*100",
    "roa": "={pretax_profit}/(({total_assets}+{total_assets@prev})/2)*100",
    "roe": "={net_profit}/{equity}*100",
    "gross_margin": "={gross_profit}/{revenue}*100",
    "net_margin": "={net_profit}/{revenue}*100",
    "abs_liquidity": "={cash}/{current_liabilities}",
    "quick_liquidity": "=({cash}+{st_investments}+{receivables})/{current_liabilities}",
    "current_liquidity": "={current_assets}/{current_liabilities}",
    "general_solvency": "={total_assets}/({current_liabilities}+{noncurrent_liabilities})",
    "debt_ratio": "=({current_liabilities}+{noncurrent_liabilities})/{total_assets}",
    "debt_to_equity": "={total_liabilities}/{equity}",
    "interest_coverage": "={operating_profit}/{financial_expenses}",
    "receivables_turnover": "={revenue}/{receivables}",
    "receivables_days": "=365/({revenue}/{receivables})",
    "dts_turnover": "={revenue}/{current_liabilities}",
    "asset_turnover": "={revenue}/{total_assets}",
}

SCORE_XL = {
    "altman_course": "=3.3*{gross_profit}/{total_assets}+1*{revenue}/{total_assets}"
                     "+0.6*{equity}/({current_liabilities}+{noncurrent_liabilities})"
                     "+1.4*{retained_earnings}/{total_assets}+1.2*{current_assets}/{total_assets}",
    "altman_original": "=1.2*({current_assets}-{current_liabilities})/{total_assets}"
                       "+1.4*{retained_earnings}/{total_assets}+3.3*{operating_profit}/{total_assets}"
                       "+0.6*{equity}/({current_liabilities}+{noncurrent_liabilities})+1*{revenue}/{total_assets}",
    "taffler": "=0.53*{pretax_profit}/{current_liabilities}"
               "+0.13*{current_assets}/({current_liabilities}+{noncurrent_liabilities})"
               "+0.18*{current_liabilities}/{total_assets}+0.16*{revenue}/{total_assets}",
    "holder": "=0.24*{ebitda}/({current_liabilities}+{noncurrent_liabilities})"
              "+0.22*({equity}+{noncurrent_liabilities})/{total_assets}"
              "+0.16*({current_assets}-{inventory})/{total_assets}"
              "-0.87*{financial_expenses}/{revenue}"
              "-0.10*IFERROR({personnel_costs}/({ebitda}+{personnel_costs}),0)",
}

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2A78D6")


def _formula(tpl: str, rows: dict, col: int) -> str:
    def sub(m):
        key, prev = m.group(1), m.group(2)
        c = col - 1 if prev and col > 2 else col
        return f"Data!{get_column_letter(c)}{rows[key]}"
    return re.sub(r"\{(\w+)(@prev)?\}", sub, tpl)


def _zone_if(model: str, cell: str) -> str:
    zones = SCORE_INFO[model]["zones"]
    (_, t1, l1), (_, t2, l2), (_, _, l3) = zones
    return f'=IFERROR(IF({cell}<{t1},"{l1}",IF({cell}<{t2},"{l2}","{l3}")),"n/a")'


def build_workbook(d: pd.DataFrame, company: str, unit_note: str) -> bytes:
    wb = Workbook()
    years = list(d.columns)

    ws = wb.active
    ws.title = "Data"
    ws["A1"] = f"{company} - input data ({unit_note})"
    ws["A1"].font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value="Indicator")
    for j, y in enumerate(years):
        ws.cell(row=2, column=2 + j, value=int(y))
    ws.cell(row=2, column=2 + len(years), value="Statement")
    rows = {}
    for i, key in enumerate(INPUTS):
        r = 3 + i
        rows[key] = r
        ws.cell(row=r, column=1, value=INPUTS[key][0])
        for j, y in enumerate(years):
            v = d.at[key, y]
            ws.cell(row=r, column=2 + j, value=None if pd.isna(v) else float(v)).number_format = "#,##0"
        ws.cell(row=r, column=2 + len(years), value=INPUTS[key][1])

    def table(ws, title, spec, info, fmt):
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        hdr = ["Indicator"] + [int(y) for y in years] + ["Formula", "Norm / interpretation"]
        for j, h in enumerate(hdr):
            c = ws.cell(row=2, column=1 + j, value=h)
            c.font, c.fill = HEAD, FILL
        return hdr

    ws2 = wb.create_sheet("Ratios")
    table(ws2, "Financial ratios (live formulas referencing sheet Data)", RATIO_XL, RATIOS, "0.00")
    for i, (key, tpl) in enumerate(RATIO_XL.items()):
        r = 3 + i
        label, _, formula, norm, _, unit = RATIOS[key]
        ws2.cell(row=r, column=1, value=f"{label} ({unit})")
        for j in range(len(years)):
            col = 2 + j
            f = _formula(tpl, rows, col)
            if key == "roa" and j == 0:  # first year: no previous assets, use year-end (template C14)
                f = _formula("={pretax_profit}/{total_assets}*100", rows, col)
            ws2.cell(row=r, column=col, value=f"=IFERROR({f[1:]},\"n/a\")").number_format = "0.00"
        ws2.cell(row=r, column=2 + len(years), value=formula)
        ws2.cell(row=r, column=3 + len(years), value=norm)

    ws3 = wb.create_sheet("Bankruptcy models")
    table(ws3, "Bankruptcy-risk scores (Solvency.xlsx models)", SCORE_XL, SCORE_INFO, "0.000")
    r = 3
    for key, tpl in SCORE_XL.items():
        ws3.cell(row=r, column=1, value=SCORE_INFO[key]["label"])
        ws3.cell(row=r + 1, column=1, value="   interpretation")
        for j in range(len(years)):
            col = 2 + j
            f = _formula(tpl, rows, col)
            ws3.cell(row=r, column=col, value=f"=IFERROR({f[1:]},\"n/a\")").number_format = "0.000"
            ws3.cell(row=r + 1, column=col, value=_zone_if(key, f"{get_column_letter(col)}{r}"))
        ws3.cell(row=r, column=2 + len(years), value=SCORE_INFO[key]["formula"].replace("\n", " | "))
        ws3.cell(row=r, column=3 + len(years),
                 value="; ".join(f"{lab}" + (f" < {hi}" if hi != float('inf') else "")
                                 for _, hi, lab in SCORE_INFO[key]["zones"]))
        r += 3

    for sh in (ws, ws2, ws3):
        sh.column_dimensions["A"].width = 48
        for j in range(len(years)):
            sh.column_dimensions[get_column_letter(2 + j)].width = 16
        sh.column_dimensions[get_column_letter(2 + len(years))].width = 60
        sh.freeze_panes = "B3"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
